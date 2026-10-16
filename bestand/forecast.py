"""Forecast (R-F1 bis R-F12): Liefertermine aus der Forecast-Liste („Forecast Cockpit“, Blatt „Auftragsübersicht“).

Nachbau von `Forecast_Importieren` (modBestand.tpl) und der Formel FORE (dashformel.py):
- Import: Blatt und Kopfzeile suchen, je Zeile AB-Schlüssel ohne Los (R-AB5) und Termin_Abgleich, alle Spalten
  mitspeichern; jeder Import ersetzt den alten Forecast ganz.
- Ansicht: ABs der Bestandsliste, die nicht Komplett sind, mit frühestem Termin, Ampel und Sortierung nach Termin.
Werte der ABs (Status, Prozent) kommen aus dem Rechenkern über ansicht.baum(); Sortierung innerhalb eines AB wie
im Dashboard (dashboard.sortieren).
"""

import io
import re
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

from django.db import transaction
from django.db.models import Min
from django.utils import timezone

from . import dashboard, rechnen
from .ab_nummer import forecast_schluessel

MAX_GROESSE = 20 * 1024 * 1024  # die Vorlage hat 1,1 MB; Platz für ein paar Jahre Wachstum
ENDUNGEN = (".xlsx", ".xlsm")  # wie beim AB-Import (Ben, 14.10.2026)
BLATT_ZUERST = "Auftragsübersicht"  # R-F2 (Bens Entscheidung F9)
GELIEFERT = "bereits geliefert"  # R-F3 (F10)

# R-F7 Ampel (Kalendertage ab heute)
ROT_TAGE, GELB_TAGE = 14, 30
ROT, GELB, WEISS = "rot", "gelb", "weiss"

# R-F8 Filter Liefertermin: (Wert in der Adresse, Beschriftung, Tage oder None = alle)
ZEITRAEUME = [("14", "Nächste 14 Tage", ROT_TAGE), ("30", "Nächste 30 Tage", GELB_TAGE), ("alle", "Alle", None)]
STANDARD_ZEITRAUM = "30"
ALT_TAGE = 8  # R-F10: „Forecast vom“ wird orange, wenn älter als 8 Tage

FEHLERWERTE = {
    "#N/A",
    "#NV",
    "#DIV/0!",
    "#VALUE!",
    "#WERT!",
    "#REF!",
    "#BEZUG!",
    "#NAME?",
    "#NUM!",
    "#ZAHL!",
    "#NULL!",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
}


class ForecastFehler(Exception):
    """Verständliche Meldung für den Benutzer."""


# ---------- R-F1 Quelle ----------
def quelle_bereinigen(q: str) -> str:
    """SharePoint-Link aus der Adresszeile in den Pfad der Datei umbauen (wie QuelleBereinigen):
    Anführungszeichen außen weg; bei http… alles ab „?“ weg und „/:x:/r/“ → „/“."""
    q = (q or "").strip()
    if len(q) > 1 and q.startswith('"') and q.endswith('"'):
        q = q[1:-1]
    if q[:4].lower() == "http":
        q = q.split("?", 1)[0]
        q = re.sub(r"/:x:/r/", "/", q, flags=re.IGNORECASE)
    return q


# ---------- Zellwerte ----------
def _wert(v):
    """Fehlerwerte (#NV usw.) gelten als leer (wie IsError → Empty)."""
    if isinstance(v, str) and v.strip().upper() in FEHLERWERTE:
        return None
    return v


def _txt(v) -> str:
    """VBA Txt(): leer → „“, sonst Text ohne Leerzeichen vorn/hinten."""
    v = _wert(v)
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y") if v.time() == time(0) else v.strftime("%d.%m.%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, bool):
        return "WAHR" if v else "FALSCH"
    if isinstance(v, (int, float)):
        return rechnen.zahl_text(v)
    return str(v).strip()


def kopf_norm(x) -> str:
    """R-F2 Normierung: klein, ohne Leerzeichen, „.“, „-“, „_“ und Zeilenumbruch."""
    t = _txt(x).lower()
    for z in (" ", ".", "-", "_", "\n", "\r"):
        t = t.replace(z, "")
    return t


_DATUM_DE = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4}|\d{2})(?:[ T]+\d{1,2}:\d{2}(?::\d{2})?)?$")
_DATUM_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]+\d{1,2}:\d{2}(?::\d{2})?)?$")
_DATUM_SCHRAEG = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


def datum(x) -> date | None:
    """R-F3 Datum erkennen: echtes Datum, Zahl zwischen 20000 und 80000 (Excel-Seriennummer) oder Datumstext
    („13.10.2026“, „13.10.26“, „2026-10-13“, auch mit Uhrzeit); die Uhrzeit wird abgeschnitten."""
    x = _wert(x)
    if x is None or isinstance(x, bool):
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    if isinstance(x, (int, float)):
        if 20000 < x < 80000:
            return date(1899, 12, 30) + timedelta(days=int(x))
        return None
    t = str(x).strip()
    try:
        if m := _DATUM_DE.match(t):
            tag, monat, jahr = int(m[1]), int(m[2]), int(m[3])
            if len(m[3]) == 2:  # wie Excel: 00–29 → 20xx, 30–99 → 19xx
                jahr += 2000 if jahr < 30 else 1900
            return date(jahr, monat, tag)
        if m := _DATUM_ISO.match(t):
            return date(int(m[1]), int(m[2]), int(m[3]))
        if m := _DATUM_SCHRAEG.match(t):
            return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None
    return None


# ---------- R-F2 bis R-F4 Datei lesen ----------
@dataclass
class Gelesen:
    datei: str
    blatt: str
    stand: str
    kopf: list
    zeilen: list = field(default_factory=list)  # (Schlüssel, Termin oder None, [Werte als Text])

    @property
    def abs_mit_termin(self) -> set:
        return {k for k, t, _ in self.zeilen if t}


def _kopfzeile(zeilen) -> tuple[int, int, int, int] | None:
    """R-F2: in den Zeilen 1–40, Spalten 1–80 die erste Zeile mit „Auftrag.-Nr“ (auftragnr/auftragsnr) und
    „Liefertermin“; „Ampel“ ist freiwillig. Liefert (Zeile, Spalte AB, Spalte Termin, Spalte Ampel), 1-basiert."""
    for r, zeile in enumerate(zeilen[:40], start=1):
        s_ab = s_t = s_amp = 0
        for c, v in enumerate(zeile[:80], start=1):
            t = kopf_norm(v)
            if t in ("auftragnr", "auftragsnr"):
                s_ab = s_ab or c
            elif t == "liefertermin":
                s_t = s_t or c
            elif t == "ampel":
                s_amp = s_amp or c
        if s_ab and s_t:
            return r, s_ab, s_t, s_amp
    return None


def _blatt_lesen(ws):
    """Alle Zeilen eines Blatts als Listen von Werten (nur lesen)."""
    return [list(z) for z in ws.iter_rows(values_only=True)]


def lesen(inhalt: bytes, dateiname: str) -> Gelesen:
    """Eine Forecast-Liste prüfen und lesen. Makros werden nie ausgeführt (openpyxl liest nur Werte)."""
    name = (dateiname or "").lower()
    if name.endswith(".xls"):
        raise ForecastFehler("Alte .xls-Dateien werden nicht gelesen. Bitte in Excel als .xlsx oder .xlsm speichern.")
    if not name.endswith(ENDUNGEN):
        raise ForecastFehler("Nur Excel-Dateien (.xlsx oder .xlsm).")
    if len(inhalt) > MAX_GROESSE:
        raise ForecastFehler(f"Die Datei ist größer als {MAX_GROESSE // (1024 * 1024)} MB.")
    if not zipfile.is_zipfile(io.BytesIO(inhalt)):
        raise ForecastFehler("Die Datei ist keine gültige Excel-Datei.")
    import openpyxl

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(
                io.BytesIO(inhalt), read_only=True, data_only=True, keep_vba=False, keep_links=False
            )
    except Exception:
        raise ForecastFehler("Die Datei kann nicht gelesen werden (beschädigt oder kein Excel-Format).")
    try:
        blaetter = list(wb.worksheets)
        reihenfolge = [ws for ws in blaetter if ws.title == BLATT_ZUERST] + [
            ws for ws in blaetter if ws.title != BLATT_ZUERST
        ]
        for ws in reihenfolge:
            kopf40 = [list(z) for z in ws.iter_rows(min_row=1, max_row=40, max_col=80, values_only=True)]
            treffer = _kopfzeile(kopf40)
            if treffer:
                return _auswerten(ws.title, _blatt_lesen(ws), treffer, dateiname)
    finally:
        wb.close()
    raise ForecastFehler(
        f"In „{dateiname}“ wurde keine Kopfzeile mit „Auftrag.-Nr“ und „Liefertermin“ gefunden. "
        f"Gesucht wird zuerst im Blatt „{BLATT_ZUERST}“."
    )


def _auswerten(blatt: str, zeilen: list, treffer, dateiname: str) -> Gelesen:
    kz, s_ab, s_t, s_amp = treffer
    kopfzeile = zeilen[kz - 1]
    letzte = max((c for c, v in enumerate(kopfzeile, start=1) if _txt(v)), default=0)  # wie End(xlToLeft)

    def zelle(z, c):
        return z[c - 1] if c - 1 < len(z) else None

    # R-F4 Stand: Zelle oberhalb der Kopfzeile, die mit „Stand:“ beginnt
    stand = ""
    for z in zeilen[: kz - 1]:
        for c in range(1, letzte + 1):
            t = _txt(zelle(z, c))
            if t[:6].lower() == "stand:":
                stand = t
                break
        if stand:
            break

    kopf = [_txt(zelle(kopfzeile, c)) or f"Spalte {c}" for c in range(1, letzte + 1)]
    g = Gelesen(datei=dateiname, blatt=blatt, stand=stand, kopf=kopf)
    for z in zeilen[kz:]:
        k = forecast_schluessel(_txt(zelle(z, s_ab)))  # R-F3: nur Zeilen mit Schlüssel (R-AB5)
        if not k:
            continue
        if s_amp and _txt(zelle(z, s_amp)).lower() == GELIEFERT:
            termin = None  # bereits geliefert → kein Termin
        else:
            termin = datum(zelle(z, s_t))
        g.zeilen.append((k, termin, [_txt(zelle(z, c)) for c in range(1, letzte + 1)]))
    if not g.zeilen:
        raise ForecastFehler("Die Forecast-Liste enthält keine Auftragsnummern.")
    return g


@dataclass
class Ergebnis:
    zeilen: int
    abs_mit_termin: int
    treffer: int
    abs_bestand: int

    @property
    def meldung(self) -> str:
        """R-F5."""
        return (
            f"Forecast importiert: {self.zeilen} Zeilen, {self.abs_mit_termin} ABs mit offenem Liefertermin; "
            f"{self.treffer} von {self.abs_bestand} ABs der Bestandsliste haben einen Liefertermin."
        )


@transaction.atomic
def speichern(g: Gelesen, benutzer=None) -> Ergebnis:
    """R-F4: ersetzt den ganzen Forecast. Liefert die Zahlen für die Meldung (R-F5)."""
    from .models import Auftrag, ForecastImport, ForecastZeile

    ForecastImport.objects.all().delete()
    f = ForecastImport.objects.create(
        zeit=timezone.now(),
        datei=g.datei[:255],
        blatt=g.blatt[:120],
        stand=g.stand[:200],
        kopf=g.kopf,
        benutzer=benutzer if getattr(benutzer, "is_authenticated", False) else None,
    )
    ForecastZeile.objects.bulk_create(
        [ForecastZeile(forecast=f, nr=i, schluessel=k, termin=t, werte=w) for i, (k, t, w) in enumerate(g.zeilen, 1)],
        batch_size=500,
    )
    mit = g.abs_mit_termin
    abs_ = list(Auftrag.objects.filter(bereich="bestand").values_list("forecast_schluessel", "ab_anzeige"))
    treffer = sum((fk or forecast_schluessel(anz)) in mit for fk, anz in abs_)
    return Ergebnis(len(g.zeilen), len(mit), treffer, len(abs_))


def aktuell():
    from .models import ForecastImport

    return ForecastImport.objects.order_by("-zeit").first()


def termine() -> dict:
    """R-F6: je AB-Schlüssel (ohne Los) der früheste Termin_Abgleich; Zeilen ohne Termin zählen nicht."""
    from .models import ForecastZeile

    return {
        z["schluessel"]: z["t"]
        for z in ForecastZeile.objects.filter(termin__isnull=False).values("schluessel").annotate(t=Min("termin"))
    }


def termin_von(a, termine_: dict) -> date | None:
    """Termin eines AB (auch eines Loses: dann der des Haupt-AB, R-AB5)."""
    return termine_.get(a.o.forecast_schluessel or forecast_schluessel(a.o.ab_anzeige))


def ampel(termin: date | None, heute: date) -> str:
    """R-F7: bis heute + 14 (auch überfällig) rot, bis heute + 30 gelb, sonst oder ohne Termin weiß."""
    if termin is None:
        return WEISS
    tage = (termin - heute).days
    return ROT if tage <= ROT_TAGE else GELB if tage <= GELB_TAGE else WEISS


# ---------- W11: entnommene ABs bleiben mit Status „Schweißen“ sichtbar, bis sie archiviert sind ----------
def mit_schweissen(abs_: list) -> list:
    """Die ABs der Bestandsliste plus die ganzen ABs im Schweißen (ABZ mit schweissen=True). Ist ein AB gerade wieder in
    der Bestandsliste (neu importiert oder zurückgeholt), zählt nur dieser."""
    from . import schweissen

    vorhanden = {a.o.ab_schluessel for a in abs_}
    return list(abs_) + [a for a in schweissen.forecast_abs() if a.o.ab_schluessel not in vorhanden]


# ---------- R-F8 bis R-F10 Ansicht ----------
def liste(
    abs_: list, termine_: dict, heute: date, zeitraum: str = STANDARD_ZEITRAUM, suche: str = "", kunde: str = ""
) -> list:
    """R-F8: nur ABs mit Status ≠ Komplett, mit allen BGs und Teilen; Filter AB-Suche, Kunde, Liefertermin.
    R-F9: Termin aufsteigend (ohne Termin unten), dann AB-Fortschritt aufsteigend, dann AB-Nummer;
    innerhalb des AB wie im Dashboard (R-DB3)."""
    tage = {w: t for w, _, t in ZEITRAEUME}.get(zeitraum, GELB_TAGE)
    auswahl = []
    for a in dashboard.filtern(abs_, suche, kunde):
        if a.w.status == rechnen.KOMPLETT and not a.schweissen:  # W11: im Schweißen bleibt er sichtbar
            continue
        t = termin_von(a, termine_)
        if tage is not None and (t is None or (t - heute).days > tage):  # 14/30: mit Überfälligen, ohne „kein Termin“
            continue
        auswahl.append(a)
    ergebnis = dashboard.sortieren(auswahl)
    for d in ergebnis:
        d.termin = termin_von(d.a, termine_)
        d.ampel = ampel(d.termin, heute)
    ergebnis.sort(
        key=lambda d: (d.termin is None, d.termin or date.max, d.a.w.prozent, dashboard._text(d.a.o.ab_anzeige))
    )
    return ergebnis


def kopf(liste_: list, f, heute: date) -> dict:
    """R-F10: Zähler über die angezeigten ABs, „Forecast vom“ (orange, wenn älter als 8 Tage), „Heute“."""
    vom = timezone.localtime(f.zeit).date() if f else None
    return {
        "rot": sum(d.ampel == ROT for d in liste_),
        "gelb": sum(d.ampel == GELB for d in liste_),
        "ohne": sum(d.termin is None for d in liste_),
        "ab": len(liste_),
        "vom": vom,
        "alt": bool(vom) and (heute - vom).days > ALT_TAGE,
        "heute": heute,
    }


def kachel(abs_: list, heute: date) -> dict | None:
    """Kachel auf Startseite und Dashboard: rote und gelbe ABs (alle offenen, ohne Filter); None = noch kein Forecast."""
    f = aktuell()
    if not f:
        return None
    t = termine()
    farben = [
        ampel(termin_von(a, t), heute) for a in mit_schweissen(abs_) if a.w.status != rechnen.KOMPLETT or a.schweissen
    ]
    return {"rot": farben.count(ROT), "gelb": farben.count(GELB), "vom": timezone.localtime(f.zeit).date()}
