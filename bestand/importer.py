"""Import einer Kommissionierliste (R-I1 bis R-I9, R-D1 bis R-D5).

Nachbau von ImportiereBlatt, KopfWert, IstBGZeile, MerkeMengen in excel_quellen/modBestand.tpl.
Dateien werden nur gelesen (openpyxl): Makros werden nie ausgeführt, Formeln nicht berechnet
(es zählt der zuletzt in Excel gespeicherte Wert).

Ablauf: lesen() prüft eine Datei und liefert ein Gelesen (ohne Datenbank, als JSON speicherbar,
damit die Rückfrage „ersetzen?“ dazwischen passt). speichern() schreibt in die Datenbank.
"""

import io
import warnings
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.db.models import Max

from .ab_nummer import ab_form, ab_schluessel, forecast_schluessel
from .models import Auftrag, Baugruppe, Teil
from .rechnen import MAX_MENGE, daten_stempeln, zahl_text

MAX_GROESSE = 10 * 1024 * 1024  # 10 MB je Datei (die Beispiele haben unter 200 KB)
ENDUNGEN = (".xlsx", ".xlsm")  # .xls gibt es laut Ben nicht mehr (14.10.2026)

# R-D1: die 12 Spalten der Vorlage → Feldname im Modell
SPALTEN = [
    ("Zeichnungsnummer", "zeichnungsnummer"),
    ("Benennung", "benennung"),
    ("Menge / BG", "menge_bg"),
    ("Menge Gesamt", "menge_gesamt"),
    ("Index", "index"),
    ("LK Nummer", "lk_nummer"),
    ("Arbeitsgang", "arbeitsgang"),
    ("Material", "material"),
    ("Profil", "profil"),
    ("Abmaße", "abmasse"),
    ("Vorhanden", "vorhanden"),
    ("Bemerkung", "bemerkung"),
]
# R-D2 / R-I4: Beschriftungen im Kopf der Liste
KOPF = [
    ("Fertigungs AB", "fertigungs_ab"),
    ("Haupt AB", "haupt_ab"),
    ("Dispoliste", "dispoliste"),
    ("Ersteller", "ersteller"),
    ("Datum", "datum"),
    ("Verkaufsartikel", "verkaufsartikel"),
    ("Kunde", "kunde"),
]
BG_FARBE = "1F497D"  # RGB(31,73,125), Dunkelblau der Baugruppen
OHNE_BG = "(ohne BG)"


class ImportFehler(Exception):
    """Datei kann nicht übernommen werden; Text ist für den Benutzer."""


@dataclass
class Gelesen:
    datei: str
    kopf: dict
    zeilen: list = field(default_factory=list)  # {"ebene": "BG"|"Teil", "bg": Name, "werte": {...}}
    hinweise: list = field(default_factory=list)

    @property
    def ab(self) -> str:
        return ab_form(self.kopf.get("fertigungs_ab", ""))

    @property
    def anz_bg(self):
        return sum(z["ebene"] == "BG" for z in self.zeilen)

    @property
    def anz_teile(self):
        return sum(z["ebene"] == "Teil" for z in self.zeilen)

    def als_dict(self):
        return asdict(self)

    @classmethod
    def aus_dict(cls, d):
        return cls(**d)


# ---------- Zellwerte ----------
def _txt(v) -> str:
    """VBA Txt(): Fehler/leer → „“, sonst Text getrimmt. Zahlen ohne „,0“ (Text bleibt Text, R-I7)."""
    if v is None or isinstance(v, bool):
        return "" if v is None else str(v)
    if isinstance(v, (int, float, Decimal)):
        return zahl_text(v)
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    return str(v).strip()


def _zahl(v):
    """Zahl oder als Text geschriebene Zahl („21“, „0,5“) → Decimal; sonst None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        d = Decimal(str(v))
        return d if d.is_finite() else None
    if isinstance(v, str):
        t = v.strip().replace(",", ".")
        if not t:
            return None
        try:
            d = Decimal(t)
        except InvalidOperation:
            return None
        return d if d.is_finite() else None
    return None


def _wert(zelle):
    """Wert einer Zelle; Fehlerwerte (#NV usw.) werden leer (R-I7)."""
    return None if zelle.data_type == "e" else zelle.value


def _ist_bg_zelle(zelle) -> bool:
    """R-I6, Regel 1: Füllung 1F497D, oder irgendeine Füllung mit fetter weißer Schrift."""
    f = zelle.fill
    if f is None or f.fill_type in (None, "none"):
        return False
    farbe = f.fgColor if f.fill_type == "solid" else f.start_color
    if farbe is not None and farbe.type == "rgb" and str(farbe.rgb).upper().endswith(BG_FARBE):
        return True
    schrift = zelle.font
    if schrift is not None and schrift.b and schrift.color is not None:
        c = schrift.color
        weiss = (c.type == "rgb" and str(c.rgb).upper().endswith("FFFFFF")) or (c.type == "theme" and c.theme == 0)
        return bool(weiss)
    return False


# ---------- Datei lesen ----------
def _kopfzeile(ws) -> int:
    """R-I2: erste Zeile 1–60, in deren Spalten A–E eine Zelle genau „Zeichnungsnummer“ enthält."""
    for r in range(1, 61):
        for c in range(1, 6):
            if _txt(_wert(ws.cell(r, c))).lower() == "zeichnungsnummer":
                return r
    return 0


def _norm(t: str) -> str:
    return t.lower().replace(" ", "").replace("\n", "").replace("\r", "")


def _spalten(ws, kz) -> dict:
    """R-I3: Spalten per Name in der Kopfzeile (Spalten 1–40)."""
    gefunden = {}
    namen = {_norm(name): feld for name, feld in SPALTEN}
    for c in range(1, 41):
        t = _norm(_txt(_wert(ws.cell(kz, c))))
        if t in namen and namen[t] not in gefunden:
            gefunden[namen[t]] = c
    return gefunden


def _kopfwert(ws, kz, etikett) -> str:
    """R-I4: Beschriftung oberhalb der Kopfzeile (Spalten 1–15), Wert in einer der 3 Zellen rechts daneben."""
    for r in range(1, kz):
        for c in range(1, 16):
            t = _txt(_wert(ws.cell(r, c)))
            if t and t.replace(":", "").strip().lower() == etikett.lower():
                for k in range(c + 1, c + 4):
                    v = _wert(ws.cell(r, k))
                    if isinstance(v, (datetime, date)):
                        return v.strftime("%d.%m.%Y")
                    if _txt(v):
                        return _txt(v)
    return ""


def lesen(inhalt: bytes, dateiname: str) -> Gelesen:
    """Eine Kommissionierliste prüfen und lesen. Wirft ImportFehler mit verständlichem Text."""
    name = dateiname.lower()
    if name.endswith(".xls"):
        raise ImportFehler("Alte .xls-Dateien werden nicht gelesen. Bitte in Excel als .xlsx speichern.")
    if not name.endswith(ENDUNGEN):
        raise ImportFehler("Nur Excel-Dateien (.xlsx oder .xlsm).")
    if len(inhalt) > MAX_GROESSE:
        raise ImportFehler("Die Datei ist größer als 10 MB.")
    if not zipfile.is_zipfile(io.BytesIO(inhalt)):
        raise ImportFehler("Die Datei ist keine gültige Excel-Datei.")
    import openpyxl

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(inhalt), data_only=True, keep_vba=False, keep_links=False)
    except Exception:
        raise ImportFehler("Die Datei kann nicht gelesen werden (beschädigt oder kein Excel-Format).")

    ws, kz = None, 0
    for blatt in wb.worksheets:
        kz = _kopfzeile(blatt)
        if kz:
            ws = blatt
            break
    if ws is None:
        raise ImportFehler("Keine Kopfzeile mit „Zeichnungsnummer“ gefunden.")
    sp = _spalten(ws, kz)
    if "zeichnungsnummer" not in sp or "menge_gesamt" not in sp:
        raise ImportFehler(f"Spalte „Zeichnungsnummer“ oder „Menge Gesamt“ fehlt (Blatt „{ws.title}“).")

    kopf = {feld: _kopfwert(ws, kz, etikett) for etikett, feld in KOPF}
    kopf["haupt_ab"] = ab_form(kopf["haupt_ab"]) if kopf["haupt_ab"] else ""
    g = Gelesen(datei=dateiname, kopf=kopf)

    # R-I5: Datenzeilen = Zeichnungsnummer oder Benennung gefüllt
    roh = []
    for r in range(kz + 1, ws.max_row + 1):
        zn = ws.cell(r, sp["zeichnungsnummer"])
        bn = _txt(_wert(ws.cell(r, sp["benennung"]))) if "benennung" in sp else ""
        if _txt(_wert(zn)) or bn:
            roh.append((r, _ist_bg_zelle(zn)))
    if not roh:
        raise ImportFehler("Keine Daten unter der Kopfzeile gefunden.")
    # R-I6, Regel 2 (nur wenn keine BG über die Farbe): LK-Nummer ganzzahlig > 0 und durch 100 teilbar
    if not any(bg for _, bg in roh) and "lk_nummer" in sp:
        neu = []
        for r, _ in roh:
            lk = _zahl(_wert(ws.cell(r, sp["lk_nummer"])))
            neu.append((r, lk is not None and lk > 0 and lk == lk.to_integral_value() and int(lk) % 100 == 0))
        roh = neu
    # Festlegung Ben 14.10.2026 (weicht von Excel ab): ohne BG unter der Kopfzeile wird die Haupt-Baugruppe
    # direkt über der Kopfzeile gelesen, wenn dort Zeichnungsnummer und eine Zahl bei „Menge Gesamt“ stehen
    haupt = kz - 1
    if (
        not any(bg for _, bg in roh)
        and haupt >= 1
        and _txt(_wert(ws.cell(haupt, sp["zeichnungsnummer"])))
        and _zahl(_wert(ws.cell(haupt, sp["menge_gesamt"]))) is not None
    ):
        roh.insert(0, (haupt, True))
        g.hinweise.append(f"Haupt-Baugruppe aus Zeile {haupt} (über der Kopfzeile) übernommen.")

    bg_name = OHNE_BG
    zu_gross = []
    for r, ist_bg in roh:
        w = {}
        for _, feld in SPALTEN:
            v = _wert(ws.cell(r, sp[feld])) if feld in sp else None
            if feld in ("menge_gesamt", "vorhanden"):
                z = _zahl(v)
                if z is not None and abs(z) > MAX_MENGE:  # R-M4: nicht übernehmen, in der Vorschau melden
                    z = None
                    zu_gross.append(r)
                w[feld] = str(z) if z is not None else None
                if feld == "menge_gesamt" and z is None and _txt(v):
                    w["menge_gesamt_text"] = _txt(v)  # Text-Menge: zählt als 0, wird aber angezeigt
            else:
                w[feld] = _txt(v)
        if ist_bg:
            bg_name = f"{w['lk_nummer']} {w['benennung']}".strip() or f"BG Zeile {r}"  # R-D4
            w["vorhanden"] = None  # R-D3, R-I7
            g.zeilen.append({"ebene": "BG", "bg": bg_name, "werte": w})
        else:
            g.zeilen.append({"ebene": "Teil", "bg": bg_name, "werte": w})

    if zu_gross:
        zeilen_text = ", ".join(str(r) for r in sorted(set(zu_gross))[:10])
        g.hinweise.append(
            f"Zu große Zahl (über 999.999.999) in Zeile {zeilen_text}: nicht übernommen. "
            "Bei „Menge Gesamt“ steht sie als Text und zählt als 0, bei „Vorhanden“ bleibt das Feld leer. "
            "Bitte die Spalten in der Datei prüfen."
        )

    if not g.anz_bg:
        g.hinweise.append(f"Keine Baugruppe erkannt: alle {g.anz_teile} Teile stehen direkt unter dem AB.")
    schluessel = [_teil_schluessel(z["bg"], z["werte"]) for z in g.zeilen if z["ebene"] == "Teil"]
    doppelt = sorted({s for s in schluessel if schluessel.count(s) > 1})
    if doppelt:
        g.hinweise.append(
            "Doppelte Teile (gleiche BG, Zeichnungsnummer und LK): "
            + "; ".join(s.replace("|", " · ") for s in doppelt)
            + ". Beim Ersetzen bekommen beide die Menge der ersten Zeile (wie Excel)."
        )
    return g


def _teil_schluessel(bg, w) -> str:
    """R-I8: Schlüssel je Teil = BG-Name | Zeichnungsnummer | LK-Nummer (ohne Groß-/Kleinschreibung)."""
    return f"{bg}|{w['zeichnungsnummer']}|{w['lk_nummer']}".lower()


def vorhandener_ab(ab: str):
    """AB mit gleichem Schlüssel (R-AB3), auch in anderer Los-Schreibweise."""
    return Auftrag.objects.filter(ab_schluessel=ab_schluessel(ab)).first()


# ---------- in die Datenbank ----------
@transaction.atomic
def speichern(g: Gelesen, benutzer=None, heute: date | None = None, ab: str | None = None) -> tuple[Auftrag, bool]:
    """Einen gelesenen AB anlegen oder ersetzen (R-I8, R-I9). Liefert (Auftrag, ersetzt)."""
    from django.utils import timezone

    heute = heute or timezone.localdate()
    ab = ab_form(ab or g.ab)
    if not ab:
        raise ImportFehler("Fertigungs-AB-Nummer fehlt.")
    alt = Auftrag.objects.select_for_update().filter(ab_schluessel=ab_schluessel(ab)).first()
    merk = {"v": {}, "l": {}, "e": {}, "f": {}}
    if alt:
        if ab_form(alt.ab_anzeige) == alt.ab_anzeige:  # bisherige Schreibweise ist korrekt → bleibt
            ab = alt.ab_anzeige
        for t in alt.teile.select_related("baugruppe").order_by("zeile"):
            s = _teil_schluessel(t.bg_name, {"zeichnungsnummer": t.zeichnungsnummer, "lk_nummer": t.lk_nummer})
            for k, wert in (("v", t.vorhanden), ("l", t.lagerort), ("e", t.datum_eingelagert), ("f", t.datum_komplett)):
                if wert not in (None, ""):
                    merk[k].setdefault(s, wert)  # erste Zeile gewinnt (wie Excel)
        alt.baugruppen.all().delete()
        alt.teile.all().delete()
        a = alt
    else:
        a = Auftrag(reihenfolge=(Auftrag.objects.aggregate(m=Max("reihenfolge"))["m"] or 0) + 1)
    k = g.kopf
    a.ab_anzeige, a.ab_schluessel, a.forecast_schluessel = ab, ab_schluessel(ab), forecast_schluessel(ab)
    a.haupt_ab, a.kunde, a.verkaufsartikel = k["haupt_ab"], k["kunde"], k["verkaufsartikel"]
    a.ersteller, a.datum_liste, a.dispoliste = k["ersteller"], k["datum"], k["dispoliste"]
    a.lagerort = ""  # W6 = wie Excel: AB-Lagerort geht verloren
    a.quelldatei = g.datei[:255]
    if benutzer is not None and getattr(benutzer, "is_authenticated", False):
        a.importiert_von = benutzer
    a.save()

    bg = None
    for nr, z in enumerate(g.zeilen, start=1):
        w = dict(z["werte"])
        soll = Decimal(w["menge_gesamt"]) if w["menge_gesamt"] is not None else None
        felder = {f: w[f] for _, f in SPALTEN if f not in ("menge_gesamt", "vorhanden")}
        felder["menge_gesamt_text"] = w.get("menge_gesamt_text", "")
        if z["ebene"] == "BG":
            bg = Baugruppe.objects.create(
                auftrag=a, zeile=nr, menge_gesamt=soll, name=z["bg"], kurz=z["bg"].split(" ")[0], **felder
            )
            continue
        s = _teil_schluessel(z["bg"], w)
        v = Decimal(w["vorhanden"]) if w["vorhanden"] is not None else merk["v"].get(s)
        ein, fertig = daten_stempeln(v, soll, merk["e"].get(s), merk["f"].get(s), heute)  # R-T1 / R-I9
        felder["lagerort"] = merk["l"].get(s, "")
        Teil.objects.create(
            auftrag=a,
            baugruppe=bg if z["bg"] != OHNE_BG else None,
            zeile=nr,
            menge_gesamt=soll,
            vorhanden=v,
            datum_eingelagert=ein,
            datum_komplett=fertig,
            **felder,
        )
    return a, alt is not None
