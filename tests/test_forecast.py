"""Forecast (R-F1 bis R-F12): Import wie `Forecast_Importieren`, Ansicht wie die Excel-Formel FORE (dashformel.py)."""

import io
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from django.test import Client

from bestand import ansicht, dashboard, forecast, views_forecast
from bestand.models import Aenderung, Ansichtszustand, Forecasteinstellung, ForecastImport, ForecastZeile
from kern.rollen import ADMIN, KOMMISSIONIERUNG

from .hilfe import DATEN, HEUTE, LISTE_204, LISTE_LOS, benutzer, importiere
from .test_dashboard import ab_anlegen, karte

pytestmark = pytest.mark.django_db
KOPF = ["Filterhilfe", "Auftrag-ID", "Auftrag.-Nr", "Kunde", "Bezeichnung", "Ampel", "Liefertermin"]


# ---------- Hilfen ----------
def fc_datei(zeilen, kopf=KOPF, blatt="Auftragsübersicht", kopfzeile=10, stand="Stand: 13.10.2026", davor=()):
    """Forecast-Liste wie die Vorlage: Kopf in Zeile `kopfzeile`, „Stand:“ darüber. davor = weitere Blätter vorne."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, inhalt in davor:
        w = wb.create_sheet(name)
        for z in inhalt:
            w.append(z)
    ws = wb.create_sheet(blatt)
    if stand:
        ws.cell(max(kopfzeile - 7, 1), min(9, len(kopf)), stand)  # wie I3, aber innerhalb der Kopfspalten
    for c, k in enumerate(kopf, 1):
        ws.cell(kopfzeile, c, k)
    for r, z in enumerate(zeilen, kopfzeile + 1):
        for c, v in enumerate(z, 1):
            ws.cell(r, c, v)
    aus = io.BytesIO()
    wb.save(aus)
    return aus.getvalue()


def zeile(ab, termin, ampel=None, kunde="", text=""):
    return ["FA", f"{ab}|1", ab, kunde, text, ampel, termin]


def laden(zeilen, name="Forecast.xlsm", **kw):
    return forecast.speichern(forecast.lesen(fc_datei(zeilen, **kw), name))


def tage(n):
    return HEUTE + timedelta(days=n)


@pytest.fixture(autouse=True)
def festes_heute(monkeypatch):
    monkeypatch.setattr(views_forecast, "heute", lambda: HEUTE)


def html(client, url="/forecast/"):
    a = client.get(url)
    assert a.status_code == 200, a.status_code
    return a.content.decode()


def reihenfolge(seite):
    """AB-Nummern der PC-Karten in der Reihenfolge der Seite."""
    pc = seite.split('class="nur-handy')[0]
    return re.findall(r'class="ab-klapp ab-nr" data-klapp-ref="ab">([^<]+)<|<span class="ab-nr">([^<]+)<', pc)


def abs_der_seite(seite):
    return [a or b for a, b in reihenfolge(seite)]


def offen(nr, prozent_vh=0, kunde=""):
    """AB mit einem Teil (Soll 100, Vorhanden = Prozent): 0 → Offen, 1–99 → Teilweise, 100 → Komplett."""
    a = ab_anlegen(nr, ohne=[(nr + "-1", 100, prozent_vh or None)], kunde=kunde)
    a.forecast_schluessel = forecast.forecast_schluessel(nr)
    a.save()
    return a


# ---------- R-F1 Quelle ----------
@pytest.mark.parametrize(
    "ein, aus",
    [
        (
            "https://firma.sharepoint.com/:x:/r/sites/Plan/Dok/Forecast.xlsm?d=w123&csf=1&web=1",
            "https://firma.sharepoint.com/sites/Plan/Dok/Forecast.xlsm",
        ),
        ('"https://firma.sharepoint.com/:X:/R/sites/a/b.xlsm?x=1"', "https://firma.sharepoint.com/sites/a/b.xlsm"),
        ("  https://firma.sharepoint.com/sites/a/b.xlsm  ", "https://firma.sharepoint.com/sites/a/b.xlsm"),
        (r"C:\Forecast\Liste.xlsm", r"C:\Forecast\Liste.xlsm"),  # lokaler Pfad bleibt (auch mit „?“ wäre er kein Link)
        ("", ""),
    ],
)
def test_R_F1_sharepoint_link_bereinigen(ein, aus):
    assert forecast.quelle_bereinigen(ein) == aus


def test_R_F1_link_hinterlegen_nur_planung_und_bereinigt():
    c = Client()
    c.force_login(benutzer("p", ADMIN))
    c.post("/forecast/import/", {"quelle": "https://firma.sharepoint.com/:x:/r/sites/X/F.xlsm?d=1"})
    assert Forecasteinstellung.objects.get(pk=1).quelle == "https://firma.sharepoint.com/sites/X/F.xlsm"
    seite = html(c, "/forecast/import/")
    assert 'href="https://firma.sharepoint.com/sites/X/F.xlsm"' in seite and "Forecast-Datei öffnen" in seite


def test_R_F1_nur_http_links_werden_als_link_gezeigt():
    c = Client()
    c.force_login(benutzer("p", ADMIN))
    c.post("/forecast/import/", {"quelle": "javascript:alert(1)"})
    assert "Forecast-Datei öffnen" not in html(c, "/forecast/import/")


# ---------- R-F2 Blatt ----------
def test_R_F2_zuerst_auftragsuebersicht():
    andere = ("Planung Stahlbau", [["Auftrag.-Nr", "Liefertermin"], ["AB2699999", datetime(2026, 11, 1)]])
    g = forecast.lesen(fc_datei([zeile("AB2515204", datetime(2026, 12, 2))], davor=[andere]), "f.xlsm")
    assert g.blatt == "Auftragsübersicht" and [k for k, _, _ in g.zeilen] == ["AB2515204"]


def test_R_F2_sonst_erstes_blatt_mit_kopfzeile():
    leer = ("Info", [["nur Text"], ["Auftrag", "Termin"]])
    g = forecast.lesen(
        fc_datei([zeile("AB2515204", datetime(2026, 12, 2))], blatt="Planung Stahlbau", davor=[leer]), "f.xlsx"
    )
    assert g.blatt == "Planung Stahlbau"


@pytest.mark.parametrize(
    "ab_kopf, termin_kopf",
    [
        ("Auftrag.-Nr", "Liefertermin"),
        ("Auftrags-Nr.", "Liefer-termin"),
        ("AUFTRAG_NR", "LIEFERTERMIN"),
        ("Auftrag\nNr.", "Liefer termin"),
        (" Auftragsnr ", "Liefertermin "),
    ],
)
def test_R_F2_kopf_normiert(ab_kopf, termin_kopf):
    kopf = ["Filterhilfe", "Auftrag-ID", ab_kopf, "Kunde", "Bezeichnung", "Ampel", termin_kopf]
    g = forecast.lesen(fc_datei([zeile("AB2515204", datetime(2026, 12, 2))], kopf=kopf), "f.xlsx")
    assert g.zeilen[0][:2] == ("AB2515204", date(2026, 12, 2))


def test_R_F2_nur_zeilen_1_bis_40():
    assert forecast.lesen(fc_datei([zeile("AB1", datetime(2026, 12, 2))], kopfzeile=40), "f.xlsx").zeilen
    with pytest.raises(forecast.ForecastFehler, match="keine Kopfzeile mit „Auftrag.-Nr“ und „Liefertermin“"):
        forecast.lesen(fc_datei([zeile("AB1", datetime(2026, 12, 2))], kopfzeile=41), "f.xlsx")


def test_R_F2_nur_spalten_1_bis_80():
    weit = ["x"] * 79 + ["Auftrag.-Nr", "Liefertermin"]
    with pytest.raises(forecast.ForecastFehler):
        forecast.lesen(fc_datei([], kopf=weit), "f.xlsx")


def test_R_F2_ohne_spalte_ampel():
    kopf = ["Auftrag.-Nr", "Liefertermin"]
    g = forecast.lesen(fc_datei([["AB2515204", datetime(2026, 12, 2)]], kopf=kopf), "f.xlsx")
    assert g.zeilen[0][1] == date(2026, 12, 2)


@pytest.mark.parametrize("name, text", [("f.xls", ".xls"), ("f.csv", "Nur Excel-Dateien"), ("f.xlsx", "keine gültige")])
def test_R_F2_falsche_dateien(name, text):
    with pytest.raises(forecast.ForecastFehler, match=text):
        forecast.lesen(b"kein excel", name)


# ---------- R-F3 Zeilen ----------
def test_R_F3_nur_zeilen_mit_schluessel():
    g = forecast.lesen(
        fc_datei(
            [
                zeile("AB2515204", datetime(2026, 12, 2)),
                zeile("", datetime(2026, 12, 2)),
                zeile("-", None),
                zeile("Summe", None),
                zeile("2616200 / Los 2", None),
                zeile("AB2515204L1", None),
            ]
        ),
        "f.xlsx",
    )
    assert [k for k, _, _ in g.zeilen] == ["AB2515204", "AB2616200", "AB2515204"]


@pytest.mark.parametrize("ampel", ["bereits geliefert", "Bereits geliefert", "  BEREITS GELIEFERT "])
def test_R_F3_bereits_geliefert_kein_termin(ampel):
    g = forecast.lesen(fc_datei([zeile("AB2515204", datetime(2026, 10, 20), ampel)]), "f.xlsx")
    assert g.zeilen[0][1] is None
    assert g.zeilen[0][2][6] == "20.10.2026"  # die Spalte selbst bleibt zum Ansehen


@pytest.mark.parametrize(
    "wert, erwartet",
    [
        (datetime(2026, 12, 2, 13, 45), date(2026, 12, 2)),  # Uhrzeit abschneiden
        (date(2026, 12, 2), date(2026, 12, 2)),
        (46358, date(2026, 12, 2)),
        (46358.75, date(2026, 12, 2)),  # Excel-Seriennummer
        ("02.12.2026", date(2026, 12, 2)),
        ("2.12.26", date(2026, 12, 2)),
        ("02.12.2026 08:30", date(2026, 12, 2)),
        ("2026-12-02", date(2026, 12, 2)),
        ("2026-12-02 08:30:00", date(2026, 12, 2)),
        (20000, None),
        (80000, None),
        (100, None),
        ("KW 49", None),
        ("31.02.2026", None),
        ("", None),
        (None, None),
        ("#N/A", None),
    ],
)
def test_R_F3_datum_erkennen(wert, erwartet):
    assert forecast.datum(wert) == erwartet


# ---------- R-F4 Stand, Speichern, Ersetzen ----------
def test_R_F4_stand_datei_blatt_und_alle_spalten():
    kopf = KOPF + ["", "Showstopper"]
    laden(
        [zeile("AB2515204", datetime(2026, 12, 2), "Gelb", "KNDS", "HALTER") + [None, "x"]],
        kopf=kopf,
        name="Forecast_Kundenauftragsliste_SCRIPT_v1.5.xlsm",
    )
    f = ForecastImport.objects.get()
    assert (f.datei, f.blatt, f.stand) == (
        "Forecast_Kundenauftragsliste_SCRIPT_v1.5.xlsm",
        "Auftragsübersicht",
        "Stand: 13.10.2026",
    )
    assert f.kopf == KOPF + ["Spalte 8", "Showstopper"]
    z = f.zeilen.get()
    assert (z.schluessel, z.termin) == ("AB2515204", date(2026, 12, 2))
    assert z.werte == ["FA", "AB2515204|1", "AB2515204", "KNDS", "HALTER", "Gelb", "02.12.2026", "", "x"]


def test_R_F4_stand_nur_mit_doppelpunkt_ueber_der_kopfzeile():
    assert forecast.lesen(fc_datei([zeile("AB1", None)], stand="Stand 13.10."), "f.xlsx").stand == ""
    assert forecast.lesen(fc_datei([zeile("AB1", None)], stand="stand: heute"), "f.xlsx").stand == "stand: heute"


def test_R_F4_jeder_import_ersetzt_alles():
    laden([zeile("AB1", datetime(2026, 12, 2)), zeile("AB2", datetime(2026, 12, 3))], name="alt.xlsm")
    laden([zeile("AB3", datetime(2026, 12, 4))], name="neu.xlsm")
    assert ForecastImport.objects.get().datei == "neu.xlsm"
    assert list(ForecastZeile.objects.values_list("schluessel", flat=True)) == ["AB3"]


def test_R_F4_vorlage_wie_analyse():
    """Die echte Vorlage (falls im Projektordner vorhanden): 780 Positionen auf „Auftragsübersicht“, Stand 13.10.2026."""
    p = Path("/mnt/project-files/web/uebergabe/beispieldaten/Forecast_Vorlage_v1.5.xlsm")
    if not p.exists():
        pytest.skip("Vorlage nur im Projektordner (enthält Kundendaten, nicht im Repository)")
    g = forecast.lesen(p.read_bytes(), p.name)
    assert (g.blatt, g.stand, len(g.zeilen)) == ("Auftragsübersicht", "Stand: 13.10.2026", 780)


# ---------- R-F5 Meldung ----------
def test_R_F5_meldung():
    importiere(LISTE_204)
    importiere(LISTE_LOS)
    offen("AB2600001")
    e = laden(
        [
            zeile("AB2515204", datetime(2026, 12, 2)),
            zeile("AB2515204", datetime(2026, 12, 9)),
            zeile("AB2616200", datetime(2026, 10, 1), "bereits geliefert"),
            zeile("AB2699999", datetime(2026, 11, 1)),
            zeile("AB2600001", None),
        ]
    )
    assert e.meldung == (
        "Forecast importiert: 5 Zeilen, 2 ABs mit offenem Liefertermin; "
        "1 von 3 ABs der Bestandsliste haben einen Liefertermin."
    )


def test_R_F5_meldung_nach_import_auf_der_seite():
    importiere(LISTE_204)
    c = Client()
    c.force_login(benutzer("p", ADMIN))
    datei = io.BytesIO(fc_datei([zeile("AB2515204", datetime(2026, 12, 2))]))
    datei.name = "Forecast.xlsm"
    a = c.post("/forecast/import/", {"datei": datei}, follow=True)
    assert (
        "Forecast importiert: 1 Zeilen, 1 ABs mit offenem Liefertermin; 1 von 1 ABs der Bestandsliste"
        in a.content.decode()
    )


def test_R_F5_fehler_wird_gemeldet():
    c = Client()
    c.force_login(benutzer("p", ADMIN))
    datei = io.BytesIO(fc_datei([zeile("", None)]))
    datei.name = "Forecast.xlsx"
    a = c.post("/forecast/import/", {"datei": datei}, follow=True)
    assert "Die Forecast-Liste enthält keine Auftragsnummern." in a.content.decode()
    assert not ForecastImport.objects.exists()


# ---------- R-F6 Termin eines AB ----------
def test_R_F6_fruehester_termin_ohne_gelieferte():
    laden(
        [
            zeile("AB1", datetime(2026, 11, 20)),
            zeile("AB1", datetime(2026, 11, 5)),
            zeile("AB1", datetime(2026, 10, 1), "bereits geliefert"),
            zeile("AB1", None),
            zeile("AB2", datetime(2026, 9, 1), "bereits geliefert"),
        ]
    )
    assert forecast.termine() == {"AB1": date(2026, 11, 5)}  # AB2: alles geliefert → kein Termin


def test_R_F6_lose_bekommen_den_termin_des_haupt_ab():
    importiere(LISTE_LOS)  # „AB2616200 / Los 1“
    laden([zeile("AB2616200", datetime(2026, 10, 19))])
    d = forecast.liste(ansicht.baum(), forecast.termine(), HEUTE, "alle")[0]
    assert (d.a.o.ab_anzeige, d.termin) == ("AB2616200 / Los 1", date(2026, 10, 19))


def test_R_F6_kein_termin(client):
    offen("AB2600001")
    laden([zeile("AB2699999", datetime(2026, 11, 1))])
    k = karte(html(client, "/forecast/?termin=alle"), "AB2600001")
    assert '<td data-feld="termin">' in k and "kein Termin" in k.split("</tr>", 1)[0]


def test_R_F6_mehrere_termine_wie_analyse():
    """Prüfdatei ABs_mehrere_Termine.xlsx: Spalte „Termin A (ohne geliefert)“ = erwarteter Termin je AB (29 ABs)."""
    p = DATEN / "ABs_mehrere_Termine.xlsx"
    forecast.speichern(forecast.lesen(p.read_bytes(), p.name))
    erwartet = {}
    for r in openpyxl.load_workbook(p, data_only=True).active.iter_rows(min_row=2, values_only=True):
        if r[0]:
            erwartet.setdefault(r[0], None)
            if r[5]:
                erwartet[r[0]] = r[5].date()
    t = forecast.termine()
    assert len(erwartet) == 29
    assert {k: t.get(k) for k in erwartet} == erwartet


# ---------- R-F7 Ampel ----------
@pytest.mark.parametrize(
    "n, farbe",
    [(-30, "rot"), (-1, "rot"), (0, "rot"), (14, "rot"), (15, "gelb"), (30, "gelb"), (31, "weiss"), (365, "weiss")],
)
def test_R_F7_ampel_kalendertage(n, farbe):
    assert forecast.ampel(tage(n), HEUTE) == farbe


def test_R_F7_ohne_termin_weiss():
    assert forecast.ampel(None, HEUTE) == "weiss"


def test_R_F7_taeglich_neu(client, monkeypatch):
    offen("AB2600001")
    laden([zeile("AB2600001", tage(20))])
    assert "fc-gelb" in karte(html(client), "AB2600001")
    monkeypatch.setattr(views_forecast, "heute", lambda: HEUTE + timedelta(days=6))
    assert "fc-rot" in karte(html(client), "AB2600001")


# ---------- R-F8 Ansicht ----------
def test_R_F8_spalten(client):
    offen("AB2600001")
    kopf = (
        html(client, "/forecast/?termin=alle").split('<table class="tabelle kopfzeile">', 1)[1].split("</thead>", 1)[0]
    )
    assert re.findall(r"<th[^>]*>([^<]*)</th>", kopf) == [
        "",
        "Arbeitsauftrag",
        "Baugruppe",
        "Einzelteil",
        "Kunde",
        "Station",
        "Lagerort",
        "Liefertermin",
        "Fortschritt",
        "",
        "Stand",
        "Planung",
    ]


def test_R_F8_nur_abs_die_nicht_komplett_sind_mit_allen_bgs_und_teilen(client):
    from .test_pruefstatus import _pruefdaten_eintragen

    _pruefdaten_eintragen()  # AB2515204 teilweise, AB2616624 komplett
    laden([zeile("AB2515204", tage(3)), zeile("AB2616624", tage(3))])
    seite = html(client, "/forecast/?ansicht=alles")
    assert abs_der_seite(seite) == ["AB2515204"]
    k = karte(seite, "AB2515204")
    assert k.count('<tr class="dbg unter"') == 3 and k.count('class="dteil teil unter"') == 25  # auch komplette BG


def test_R_F8_liefertermin_nur_in_der_ab_zeile(client):
    importiere(LISTE_204)
    laden([zeile("AB2515204", tage(3))])
    k = karte(html(client, "/forecast/?ansicht=alles"), "AB2515204")
    ab, rest = k.split("</tr>", 1)
    assert f"<b>{tage(3):%d.%m.%Y}</b>" in ab and f"{tage(3):%d.%m.%Y}" not in rest
    assert 'data-feld="status"' not in k  # Spalte Status gibt es im Forecast nicht


def test_R_F8_filter_liefertermin():
    for nr in ("AB2600001", "AB2600002", "AB2600003", "AB2600004", "AB2600005"):
        offen(nr)
    laden(
        [
            zeile("AB2600001", tage(-5)),
            zeile("AB2600002", tage(14)),
            zeile("AB2600003", tage(30)),
            zeile("AB2600004", tage(31)),
        ]
    )  # AB2600005 ohne Termin
    nrn = lambda z: [d.a.o.ab_anzeige for d in forecast.liste(ansicht.baum(), forecast.termine(), HEUTE, z)]
    assert nrn("14") == ["AB2600001", "AB2600002"]  # Überfällige dabei
    assert nrn("30") == ["AB2600001", "AB2600002", "AB2600003"]  # ohne Termin nicht
    assert nrn("alle") == ["AB2600001", "AB2600002", "AB2600003", "AB2600004", "AB2600005"]


def test_R_F8_standard_naechste_30_tage(client):
    assert forecast.ZEITRAEUME == [("14", "Nächste 14 Tage", 14), ("30", "Nächste 30 Tage", 30), ("alle", "Alle", None)]
    offen("AB2600001")
    offen("AB2600002")
    laden([zeile("AB2600001", tage(20)), zeile("AB2600002", tage(40))])
    seite = html(client)
    assert '<option value="30" selected>Nächste 30 Tage</option>' in seite
    assert abs_der_seite(seite) == ["AB2600001"]


def test_R_F8_ab_suche_und_kunde(client):
    offen("AB2600001", kunde="KNDS")
    offen("AB2600002", kunde="Rheinmetall")
    offen("AB2611111", kunde="knds")
    laden([zeile(n, tage(5)) for n in ("AB2600001", "AB2600002", "AB2611111")])
    assert abs_der_seite(html(client, "/forecast/?q=26000")) == ["AB2600001", "AB2600002"]
    assert abs_der_seite(html(client, "/forecast/?kunde=KNDS")) == ["AB2600001", "AB2611111"]
    seite = html(client, "/forecast/")
    assert re.findall(r'<option value="([^"]*)"', seite.split('id="kunde"', 1)[1].split("</select>", 1)[0]) == [
        "Alle",
        "KNDS",
        "Rheinmetall",
    ]


def test_R_F8_ansicht_und_klappen_eigener_zustand():
    u = benutzer("l", KOMMISSIONIERUNG)
    c = Client()
    c.force_login(u)
    offen("AB2600001")
    laden([zeile("AB2600001", tage(5))])
    c.get("/forecast/?ansicht=alles")
    a = c.post(
        "/forecast/klappen/",
        {"ansicht": "alles", "schluessel": "ab2600001", "offen": False},
        content_type="application/json",
    )
    assert a.status_code == 200
    assert Ansichtszustand.objects.get(benutzer=u, seite="forecast").klappen == {"ab2600001": False}
    assert not Ansichtszustand.objects.filter(benutzer=u, seite="dashboard").exists()
    assert 'aria-expanded="false" aria-label="AB2600001' in html(c, "/forecast/")  # gemerkte Ansicht „Alles“


# ---------- R-F9 Sortierung ----------
def test_R_F9_termin_dann_fortschritt_dann_ab_nummer(client):
    """Beispiel aus dem Plan: 01.01./0 % → 01.01./50 % → 01.02./20 % → …; ohne Termin ganz unten."""
    offen("AB2600005", 0)
    offen("AB2600004", 20)
    offen("AB2600003", 50)
    offen("AB2600002", 0)
    offen("AB2600001", 0)
    offen("AB2600006", 10)
    laden(
        [
            zeile("AB2600002", tage(3)),
            zeile("AB2600003", tage(3)),
            zeile("AB2600001", tage(3)),
            zeile("AB2600004", tage(10)),
        ]
    )  # 5 und 6 ohne Termin
    assert abs_der_seite(html(client, "/forecast/?termin=alle")) == [
        "AB2600001",
        "AB2600002",
        "AB2600003",
        "AB2600004",
        "AB2600005",
        "AB2600006",
    ]


def test_R_F9_innerhalb_des_ab_wie_dashboard():
    from .test_dashboard import zeilen
    from .test_pruefstatus import _pruefdaten_eintragen

    _pruefdaten_eintragen()
    laden([zeile("AB2515204", tage(3))])
    fc = forecast.liste(ansicht.baum(), forecast.termine(), HEUTE, "alle")
    db = [d for d in dashboard.sortieren(ansicht.baum()) if d.a.o.ab_anzeige == "AB2515204"]
    assert zeilen(fc) == zeilen(db)


# ---------- R-F10 Farben und Kopf ----------
def test_R_F10_ab_kraeftig_bg_und_teile_hell_balken_in_statusfarbe(client):
    importiere(LISTE_204)
    offen("AB2600002", 30)
    laden([zeile("AB2515204", tage(10)), zeile("AB2600002", tage(20))])
    seite = html(client, "/forecast/?ansicht=alles")
    assert '<table class="tabelle dash fc fc-rot">' in karte(seite, "AB2515204")
    assert '<table class="tabelle dash fc fc-gelb">' in karte(seite, "AB2600002")
    assert 'class="balken teilweise"' in karte(seite, "AB2600002")
    css = Path("kern/static/kern/css/listen.css").read_text(encoding="utf-8")
    assert ".tabelle.fc-rot tr.dab td { background: var(--ampel-rot);" in css
    assert ".tabelle.fc-gelb tr.dab td { background: var(--ampel-gelb);" in css
    assert ".tabelle.fc-rot tr.dbg td, .tabelle.fc-rot tr.dteil td { background: var(--ampel-rot-bg); }" in css
    assert ".tabelle.fc-gelb tr.dbg td, .tabelle.fc-gelb tr.dteil td { background: var(--ampel-gelb-bg); }" in css
    tokens = Path("kern/static/kern/design/tokens.css").read_text(encoding="utf-8").split('[data-theme="light"]')[1]
    for name, wert in [
        ("ampel-rot", "#ff9b9b"),
        ("ampel-gelb", "#ffe678"),
        ("ampel-rot-bg", "#ffe4e4"),
        ("ampel-gelb-bg", "#fff7d6"),
    ]:
        assert f"--{name}: {wert};" in tokens


def test_R_F10_zaehler_ueber_die_angezeigten_abs(client):
    for nr in ("AB2600001", "AB2600002", "AB2600003", "AB2600004", "AB2600005"):
        offen(nr)
    laden(
        [
            zeile("AB2600001", tage(-2)),
            zeile("AB2600002", tage(10)),
            zeile("AB2600003", tage(20)),
            zeile("AB2600004", tage(60)),
        ]
    )

    def kopf(url):
        s = html(client, url)
        return {k: int(v) for k, v in re.findall(r'data-kopf="(\w+)">(\d+)<', s)}

    assert kopf("/forecast/?termin=alle") == {"rot": 2, "gelb": 1, "ohne": 1}
    assert kopf("/forecast/?termin=14") == {"rot": 2, "gelb": 0, "ohne": 0}


@pytest.mark.parametrize("alter, orange", [(0, False), (8, False), (9, True)])
def test_R_F10_forecast_vom_orange_wenn_aelter_als_8_tage(client, alter, orange):
    laden([zeile("AB1", None)])
    ForecastImport.objects.update(
        zeit=datetime(2026, 10, 14, 10, 0, tzinfo=forecast.timezone.get_current_timezone()) - timedelta(days=alter)
    )
    seite = html(client)
    assert ('class="kachel fc-alt" data-kopf="vom"' in seite) == orange
    assert f"{(HEUTE - timedelta(days=alter)):%d.%m.%Y}" in seite.split('data-kopf="vom"', 1)[1].split("</div>", 1)[0]
    assert ">Heute<" not in seite and 'data-kopf="ab"' not in seite  # Kacheln ABs und Heute entfernt (Ben, 14.10.)
    assert 'class="knopf such-knopf"' in seite and ">BGs<" in seite and 'name="ansicht" value="bg"' in seite


def test_R_F10_noch_kein_forecast(client):
    offen("AB2600001")
    seite = html(client, "/forecast/?termin=alle")
    assert "Noch kein Forecast importiert" in seite and abs_der_seite(seite) == ["AB2600001"]


# ---------- R-F11 Rollen ----------
def test_R_F11_ansehen_darf_jeder_aber_ohne_knoepfe(client):
    seite = html(client)
    assert "Forecast importieren" not in seite and "Forecast-Daten ansehen" not in seite


@pytest.mark.parametrize("url", ["/forecast/import/", "/forecast/daten/"])
def test_R_F11_import_und_daten_nur_planung(url):
    assert Client().get(url).status_code == 302  # Gast → Anmeldung
    c = Client()
    c.force_login(benutzer("l", KOMMISSIONIERUNG))
    assert c.get(url).status_code == 403
    c.force_login(benutzer("p", ADMIN))
    assert c.get(url).status_code == 200


def test_R_F11_import_ohne_rolle_wird_abgelehnt():
    c = Client()
    c.force_login(benutzer("l", KOMMISSIONIERUNG))
    datei = io.BytesIO(fc_datei([zeile("AB1", None)]))
    datei.name = "f.xlsx"
    assert c.post("/forecast/import/", {"datei": datei}).status_code == 403
    assert not ForecastImport.objects.exists()


def test_R_F11_forecast_daten_alle_spalten_und_suche():
    laden([zeile("AB2515204", datetime(2026, 12, 2), "Gelb", "KNDS"), zeile("AB2699999", None, None, "Rheinmetall")])
    c = Client()
    c.force_login(benutzer("p", ADMIN))
    seite = html(c, "/forecast/daten/")
    for k in ["AB_Abgleich", "Termin_Abgleich"] + KOPF:
        assert f">{k}</th>" in seite
    assert "<td>KNDS</td>" in seite and "<td>Rheinmetall</td>" in seite
    seite = html(c, "/forecast/daten/?q=rhein")
    assert "<td>Rheinmetall</td>" in seite and "<td>KNDS</td>" not in seite
    assert "Forecast importieren" in html(c, "/forecast/") and "Forecast-Daten ansehen" in html(c, "/forecast/")


# ---------- R-F12 Schweißen und Forecast ----------
def test_R_F12_nur_abs_aus_der_bestandsliste():
    """Stand Phase 5. Ben, 14.10.: W11 = entnommene ABs mit Status „Schweißen“ zeigen, bis sie archiviert sind;
    das kommt mit Schweißen in Phase 6, dann wird dieser Test erweitert."""
    offen("AB2600001")
    weg = offen("AB2600002")
    weg.bereich = "schweissen"
    weg.save()
    laden([zeile("AB2600001", tage(3)), zeile("AB2600002", tage(3))])
    assert [d.a.o.ab_anzeige for d in forecast.liste(ansicht.baum(), forecast.termine(), HEUTE, "alle")] == [
        "AB2600001"
    ]


# ---------- Protokoll, Kacheln ----------
def test_forecast_import_im_protokoll():
    c = Client()
    c.force_login(benutzer("p", ADMIN))
    datei = io.BytesIO(fc_datei([zeile("AB1", None), zeile("AB2", None)]))
    datei.name = "Forecast.xlsm"
    c.post("/forecast/import/", {"datei": datei})
    e = Aenderung.objects.get()
    assert (e.art, e.objekt, e.neu, e.benutzer_name) == ("forecast", "Forecast.xlsm", "2 Zeilen", "p")


def test_forecast_kachel_nur_auf_startseite_nicht_im_dashboard(client):
    offen("AB2600001")
    offen("AB2600002")
    offen("AB2600003", 100)  # komplett: zählt nicht
    assert 'data-kachel="forecast"' not in html(client, "/dashboard/")  # Ben, 14.10.: Kachel im Dashboard entfernt
    assert "Noch kein Forecast importiert" in html(client, "/")
    laden(
        [
            zeile("AB2600001", forecast.timezone.localdate()),
            zeile("AB2600002", forecast.timezone.localdate() + timedelta(days=20)),
            zeile("AB2600003", forecast.timezone.localdate()),
        ]
    )
    for url in ("/",):
        k = html(client, url).split('data-kachel="forecast"', 1)[1].split("</a>", 1)[0]
        assert re.sub(r"<[^>]+>", "", k.split('class="zahl">', 1)[1].split("</span>\n", 1)[0]).split() == [
            "1",
            "rot",
            "·",
            "1",
            "gelb",
        ]


def test_beispiel_forecast_zeigt_alle_farben():
    from bestand.management.commands.beispieldaten import beispiel_forecast

    for p in (LISTE_204, DATEN / "Kommissionierung_Stahlbau_AB_AB2616048.xlsx", LISTE_LOS):
        importiere(p)
    forecast.speichern(forecast.lesen(beispiel_forecast(HEUTE), "b.xlsx"))
    farben = {d.a.o.ab_anzeige: d.ampel for d in forecast.liste(ansicht.baum(), forecast.termine(), HEUTE, "alle")}
    assert farben == {"AB2515204": "rot", "AB2616048": "gelb", "AB2616200 / Los 1": "weiss"}
