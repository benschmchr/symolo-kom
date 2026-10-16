"""R-I1 bis R-I9, R-D1 bis R-D5: Import einer Kommissionierliste."""

import io
from datetime import date, datetime
from decimal import Decimal as D

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

from bestand import importer
from bestand.importer import ImportFehler, lesen
from bestand.models import Auftrag, Baugruppe, Teil

pytestmark = pytest.mark.django_db

from .hilfe import HEUTE, LISTE_048, LISTE_204, LISTE_624, LISTE_LOS, importiere

KOPF = [
    "Zeichnungsnummer",
    "Benennung",
    "Menge / BG",
    "Menge Gesamt",
    "Index",
    "LK Nummer",
    "Arbeitsgang",
    "Material",
    "Profil",
    "Abmaße",
    "Vorhanden",
    "Bemerkung",
]


def xlsx(zeilen, kopf=None, kopfzeile=10, ab="AB123456", farbe=None, spalten=KOPF, extra=None):
    """Kleine Kommissionierliste bauen. zeilen: Listen mit den Werten der Spalten; farbe: {Zeilenindex: Art}."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kommissionierung"
    kopf = kopf if kopf is not None else {"Fertigungs AB:": ab, "Kunde:": "KNDS", "Datum:": datetime(2026, 7, 3)}
    for i, (k, v) in enumerate(kopf.items(), start=2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
    for c, n in enumerate(spalten, start=1):
        ws.cell(kopfzeile, c, n)
    for r, z in enumerate(zeilen, start=kopfzeile + 1):
        for c, v in enumerate(z, start=1):
            ws.cell(r, c, v)
        art = (farbe or {}).get(r - kopfzeile - 1)
        if art == "blau":
            ws.cell(r, 1).fill = PatternFill("solid", fgColor="1F497D")
        elif art == "weiss_fett":
            ws.cell(r, 1).fill = PatternFill("solid", fgColor="4472C4")
            ws.cell(r, 1).font = Font(bold=True, color="FFFFFF")
    if extra:
        extra(ws)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def z(zn, bn, soll, lk=None, vh=None):
    return [zn, bn, 1, soll, "-", lk, "", "", "", "", vh, ""]


# ---------- Beispieldateien ----------
def test_R_I_beispiel_204_zahlen_und_kopfwerte():
    g = lesen(LISTE_204.read_bytes(), LISTE_204.name)
    assert g.ab == "AB2515204"
    assert (g.anz_bg, g.anz_teile) == (3, 25)
    assert g.kopf == {
        "fertigungs_ab": "AB2515204",
        "haupt_ab": "AB2515202",
        "dispoliste": "DispositionslisteAB2515204.xlsm",
        "ersteller": "J.Steinort",
        "datum": "03.07.2026",
        "verkaufsartikel": "Halter Anzeigegeraet VO",
        "kunde": "KNDS",
    }
    assert [x["bg"] for x in g.zeilen if x["ebene"] == "BG"] == [
        "100 HALTER ANZEIGEGERAET GE",
        "200 HALTER ANZEIGEGERAET GE",
        "300 ACHSE GESCHW.",
    ]


def test_R_I6_beispiel_624_bg_mit_lk_10_ueber_farbe():
    g = lesen(LISTE_624.read_bytes(), LISTE_624.name)
    assert g.ab == "AB2616624" and g.kopf["haupt_ab"] == "-"
    assert [x["bg"] for x in g.zeilen if x["ebene"] == "BG"] == ["10 ABDECKUNG, GESCHW."]
    assert g.anz_teile == 5


def test_R_I5_haupt_baugruppe_ueber_kopfzeile_wird_baugruppe():
    # Bens Liste AB2616048: Haupt-Baugruppe steht in Zeile 9 über der Kopfzeile. Ben 14.10.2026: als BG lesen
    g = lesen(LISTE_048.read_bytes(), LISTE_048.name)
    assert any("Zeile 9" in h for h in g.hinweise) and not any("Keine Baugruppe" in h for h in g.hinweise)
    a = importiere(LISTE_048)
    assert (a.ab_anzeige, a.kunde, a.haupt_ab) == ("AB2616048", "Rheinmetall", "AB2515037")
    b = a.baugruppen.get()
    assert (b.name, b.zeichnungsnummer, b.menge_gesamt, b.kurzname) == (
        "200 Gehäuse, vollst.",
        "L41F101003-063120.000.0",
        6,
        "AB2616048_BG200",
    )
    assert a.teile.count() == 11 and all(t.baugruppe_id == b.id for t in a.teile.all())
    assert a.teile.get(zeichnungsnummer="WN2.004-3--AL5183").menge_gesamt == D("0.06")  # Soll mit Komma


def test_R_I5_haupt_baugruppe_los_1():
    a = importiere(LISTE_LOS)
    assert a.baugruppen.get().name == "10 FA - BATTERIERAHMEN Los 1/2" and a.baugruppen.get().teile.count() == 2


def test_R_I5_ab_nur_mit_teilen_ohne_baugruppe():
    g = lesen(xlsx([z("A-1", "TEIL", 2, 11), z("A-2", "TEIL", 2, 12)]), "nur_teile.xlsx")
    assert [(x["ebene"], x["bg"]) for x in g.zeilen] == [("Teil", "(ohne BG)"), ("Teil", "(ohne BG)")]
    assert any("Keine Baugruppe" in h for h in g.hinweise)


def test_R_I5_zeile_ueber_kopfzeile_nur_ohne_andere_bg_und_mit_menge():
    ueber = lambda menge: lambda ws: [ws.cell(9, c, v) for c, v in enumerate(["H-0", "HAUPT", 1, menge], 1)]
    g = lesen(xlsx([z("A-10", "BG", 1, 10), z("A-11", "T", 1, 11)], farbe={0: "blau"}, extra=ueber(3)), "x.xlsx")
    assert [x["bg"] for x in g.zeilen if x["ebene"] == "BG"] == ["10 BG"]  # farbige BG hat Vorrang
    g = lesen(xlsx([z("A-11", "T", 1, 11)], extra=ueber("Text")), "x.xlsx")
    assert g.anz_bg == 0  # keine Menge → keine BG


def test_R_AB6_los_ist_eigener_ab():
    a = importiere(LISTE_LOS)
    assert (a.ab_anzeige, a.ab_schluessel, a.forecast_schluessel) == ("AB2616200 / Los 1", "AB2616200|L1", "AB2616200")
    assert a.haupt_ab == "AB2616199" and a.teile.count() == 2
    b = importiere(LISTE_204)
    assert Auftrag.objects.count() == 2 and b.id != a.id


def test_R_D2_kopfdaten_als_eigene_felder():
    a = importiere(LISTE_204)
    assert (a.verkaufsartikel, a.kunde, a.ersteller, a.datum_liste, a.dispoliste) == (
        "Halter Anzeigegeraet VO",
        "KNDS",
        "J.Steinort",
        "03.07.2026",
        "DispositionslisteAB2515204.xlsm",
    )


def test_R_D3_R_D4_R_D5_baugruppen():
    a = importiere(LISTE_204)
    b = a.baugruppen.get(kurz="100")
    assert b.name == "100 HALTER ANZEIGEGERAET GE" and b.menge_gesamt == 21
    assert b.kurzname == "AB2515204_BG100"
    assert not hasattr(b, "vorhanden")  # BG hat keine Menge
    assert b.teile.count() == 8


def test_R_D1_alle_12_spalten_und_U2_trimmen():
    a = importiere(LISTE_204)
    t = a.teile.get(zeichnungsnummer="SK1990-977451.009.0")
    assert (
        t.benennung,
        t.menge_bg,
        t.menge_gesamt,
        t.index,
        t.lk_nummer,
        t.arbeitsgang,
        t.material,
        t.profil,
        t.abmasse,
    ) == ("HALTER", "1", 21, "-", "101", "Laser-Kantteile", "S355J2C+C (1.0579)", "Tafel", "226 x 4 x 1200")
    assert a.teile.get(zeichnungsnummer="SK1300-060201.178.0").benennung == "MASSEBUTZEN M8X10       SW22 (EDELSTAHL)"


# ---------- künstliche Dateien ----------
def test_R_I6_ersatzregel_ohne_farbe():
    g = lesen(
        xlsx(
            [
                z("A-100", "BAUGRUPPE EINS", 2, 100),
                z("A-101", "TEIL", 4, 101),
                z("A-200", "BAUGRUPPE ZWEI", 2, "200"),
                z("A-201", "TEIL", 4, 201),
            ]
        ),
        "ohne.xlsx",
    )
    assert [x["bg"] for x in g.zeilen if x["ebene"] == "BG"] == ["100 BAUGRUPPE EINS", "200 BAUGRUPPE ZWEI"]


def test_R_I6_farbe_hat_vorrang_vor_ersatzregel():
    g = lesen(
        xlsx([z("A-10", "BG ZEHN", 1, 10), z("A-11", "T", 1, 11), z("A-100", "KEINE BG", 1, 100)], farbe={0: "blau"}),
        "farbe.xlsx",
    )
    assert [x["ebene"] for x in g.zeilen] == ["BG", "Teil", "Teil"]


def test_R_I6_fuellung_mit_weisser_fetter_schrift():
    g = lesen(xlsx([z("A-10", "BG", 1, 10), z("A-11", "T", 1, 11)], farbe={0: "weiss_fett"}), "w.xlsx")
    assert g.zeilen[0]["ebene"] == "BG"


def test_R_D4_teile_vor_erster_bg_und_leerer_bg_name():
    g = lesen(
        xlsx(
            [
                z("X-1", "VORAB", 3),
                [None, None, 1, 2, "", None, "", "", "", "", None, "nur Bem."],
                z("", "TEIL NACH BG", 1),
            ],
            farbe={1: "blau"},
        ),
        "x.xlsx",
    )
    # Zeile ohne Zeichnungsnummer und Benennung wird nicht gelesen (R-I5), also auch keine BG
    assert [(x["ebene"], x["bg"]) for x in g.zeilen] == [("Teil", "(ohne BG)"), ("Teil", "(ohne BG)")]
    g = lesen(
        xlsx(
            [z("X-1", "VORAB", 3), ["B-1", None, 1, 2, "", None, "", "", "", "", None, ""], z("T-1", "T", 1)],
            farbe={1: "blau"},
        ),
        "x.xlsx",
    )
    assert [x["bg"] for x in g.zeilen] == ["(ohne BG)", "BG Zeile 12", "BG Zeile 12"]


def test_R_I2_kopfzeile_suchen_und_R_I3_pflichtspalten():
    with pytest.raises(ImportFehler, match="Zeichnungsnummer"):
        lesen(xlsx([z("A", "B", 1)], kopfzeile=61), "tief.xlsx")
    g = lesen(xlsx([z("A", "B", 1)], kopfzeile=60), "ok.xlsx")
    assert g.anz_teile == 1
    with pytest.raises(ImportFehler, match="Menge Gesamt"):
        lesen(xlsx([["A", "B"]], spalten=["Zeichnungsnummer", "Benennung"]), "ohne_soll.xlsx")


def test_R_I3_spaltennamen_ohne_leerzeichen_und_umbruch():
    spalten = [
        "ZEICHNUNGSNUMMER",
        "Benennung",
        "Menge/BG",
        "Menge\nGesamt",
        "Index",
        "LKNummer",
        "Arbeitsgang",
        "Material",
        "Profil",
        "Abmaße",
        "vorhanden",
        "Bemerkung",
    ]
    g = lesen(xlsx([z("A", "B", 7, 5, 3)], spalten=spalten), "s.xlsx")
    w = g.zeilen[0]["werte"]
    assert (w["menge_gesamt"], w["lk_nummer"], w["vorhanden"]) == ("7", "5", "3")


def test_R_I4_kopfwerte_doppelpunkt_egal_und_datum():
    def kopf(ws):
        ws["D2"] = "Verkaufsartikel"
        ws["F2"] = "  Rahmen  "  # zweite Zelle rechts daneben

    g = lesen(xlsx([z("A", "B", 1)], extra=kopf), "k.xlsx")
    assert g.kopf["verkaufsartikel"] == "Rahmen" and g.kopf["datum"] == "03.07.2026"


def test_R_I4_ohne_fertigungs_ab_wird_gefragt():
    g = lesen(xlsx([z("A", "B", 1)], kopf={"Kunde:": "KNDS"}), "ohne_ab.xlsx")
    assert g.ab == ""
    with pytest.raises(ImportFehler):
        importer.speichern(g, heute=HEUTE)


def test_R_I7_werte():
    def fehler(ws):
        ws["D12"] = "#N/A"
        ws["D12"].data_type = "e"

    g = lesen(
        xlsx(
            [
                ["007", "TEIL", 1, 4, "", 1, "", "", "", "", "viele", ""],
                [12345, "ZAHL", 1, "#NV", "", 2, "", "", "", "", " 3 ", ""],
                ["C", "BG", 1, 2, "", 100, "", "", "", "", 9, ""],
            ],
            farbe={2: "blau"},
            extra=fehler,
        ),
        "w.xlsx",
    )
    a, b, c = (x["werte"] for x in g.zeilen)
    assert a["zeichnungsnummer"] == "007" and a["vorhanden"] is None  # Text bleibt Text, Text-Menge leer
    assert b["zeichnungsnummer"] == "12345" and b["menge_gesamt"] is None and b["vorhanden"] == "3"
    assert c["vorhanden"] is None  # BG: Vorhanden ignoriert


@pytest.mark.parametrize(
    "name, inhalt, text",
    [
        ("alt.xls", b"\xd0\xcf\x11\xe0", ".xls"),
        ("liste.csv", b"a;b", "Nur Excel"),
        ("kaputt.xlsx", b"kein zip", "keine gültige"),
        ("leer.xlsm", b"PK\x05\x06" + b"\0" * 18, "nicht gelesen"),
    ],
)
def test_R_I1_nur_excel_dateien(name, inhalt, text):
    with pytest.raises(ImportFehler, match=text):
        lesen(inhalt, name)


def test_R_I1_groessenlimit():
    with pytest.raises(ImportFehler, match="10 MB"):
        lesen(b"PK" + b"0" * (importer.MAX_GROESSE + 1), "gross.xlsx")


def test_R_I1_makros_werden_nie_ausgefuehrt():
    # openpyxl liest nur Zellwerte; mit keep_vba=False wird kein Makro-Code übernommen oder ausgeführt
    import inspect

    quelle = inspect.getsource(importer.lesen)
    assert "keep_vba=False" in quelle and "data_only=True" in quelle


# ---------- Ersetzen ----------
def test_R_I8_ersetzen_uebernimmt_mengen_lagerort_daten():
    a = importiere(LISTE_204, heute=date(2026, 10, 12))
    t = a.teile.get(zeichnungsnummer="SK1990-977451.009.0")
    t.vorhanden, t.lagerort, t.datum_eingelagert, t.datum_komplett = 21, "L3", date(2026, 10, 12), date(2026, 10, 13)
    t.save()
    Auftrag.objects.filter(id=a.id).update(lagerort="L3")
    Baugruppe.objects.filter(auftrag=a).update(lagerort="L3")
    b, ersetzt = importer.speichern(lesen(LISTE_204.read_bytes(), "x.xlsx"), heute=HEUTE)
    assert ersetzt and b.id == a.id and Auftrag.objects.count() == 1
    t = b.teile.get(zeichnungsnummer="SK1990-977451.009.0")
    assert (t.vorhanden, t.lagerort, t.datum_eingelagert, t.datum_komplett) == (
        21,
        "L3",
        date(2026, 10, 12),
        date(2026, 10, 13),
    )
    assert b.lagerort == "" and set(b.baugruppen.values_list("lagerort", flat=True)) == {""}  # W6 = wie Excel
    assert b.teile.count() == 25


def test_R_I8_menge_aus_datei_hat_vorrang():
    importiere(xlsx([z("T-1", "TEIL", 5, 1)]))
    Teil.objects.update(vorhanden=2)
    a = importiere(xlsx([z("T-1", "TEIL", 5, 1, 4)]))
    assert a.teile.get().vorhanden == 4 and a.teile.get().datum_eingelagert == HEUTE


def test_R_I8_schluessel_bg_zeichnung_lk_ohne_grossschreibung():
    importiere(xlsx([z("b-1", "BG", 1, 100), z("t-1", "TEIL", 5, 101), z("T-2", "TEIL", 5, 102)], farbe={0: "blau"}))
    Teil.objects.update(vorhanden=3)
    a = importiere(
        xlsx([z("B-1", "BG", 1, 100), z("T-1", "TEIL", 5, 101), z("T-2", "TEIL", 5, 999)], farbe={0: "blau"})
    )
    assert dict(a.teile.values_list("zeichnungsnummer", "vorhanden")) == {"T-1": 3, "T-2": None}


def test_R_I8_doppelte_teile_erste_menge_gilt_und_hinweis():
    importiere(xlsx([z("T-1", "A", 5, 1), z("T-1", "B", 5, 1)]))
    Teil.objects.filter(benennung="A").update(vorhanden=1)
    Teil.objects.filter(benennung="B").update(vorhanden=2)
    g = lesen(xlsx([z("T-1", "A", 5, 1), z("T-1", "B", 5, 1)]), "d.xlsx")
    assert any("Doppelte Teile" in h for h in g.hinweise)
    a, _ = importer.speichern(g, heute=HEUTE)
    assert list(a.teile.values_list("vorhanden", flat=True)) == [1, 1]


def test_R_I8_schreibweise_und_U3_platz_bleibt():
    erster = importiere(xlsx([z("T", "T", 1)], ab="2616200 / Los 1"))
    importiere(xlsx([z("T", "T", 1)], ab="AB999"))
    a = importiere(xlsx([z("T", "T", 1)], ab="AB2616200 L01"))
    assert a.id == erster.id and a.ab_anzeige == "AB2616200 / Los 1"  # korrekte Schreibweise bleibt
    assert list(Auftrag.objects.values_list("ab_anzeige", flat=True)) == ["AB2616200 / Los 1", "AB999"]


def test_R_I9_daten_der_mengen_aus_der_datei():
    a = importiere(xlsx([z("T-1", "T", 5, 1, 5), z("T-2", "T", 5, 2, 2), z("T-3", "T", 5, 3)]))
    assert list(a.teile.values_list("datum_eingelagert", "datum_komplett")) == [
        (HEUTE, HEUTE),
        (HEUTE, None),
        (None, None),
    ]


# ---------- Obergrenze (R-M4) ----------
def test_R_M4_import_zu_grosse_zahl_wird_text_und_warnt():
    g = lesen(xlsx([z("A1", "x", 10**12, vh=2), z("A2", "y", 999_999_999, vh=10**12), z("A3", "z", 5)]), "t.xlsx")
    w = {r["werte"]["zeichnungsnummer"]: r["werte"] for r in g.zeilen if r["ebene"] == "Teil"}
    assert w["A1"]["menge_gesamt"] is None and w["A1"]["menge_gesamt_text"] == "1000000000000"
    assert w["A2"]["menge_gesamt"] == "999999999" and w["A2"]["vorhanden"] is None
    assert any("999.999.999" in h for h in g.hinweise)


def test_R_M4_import_zu_grosse_zahl_legt_seiten_nicht_lahm(client):
    from kern.rollen import KOMMISSIONIERUNG

    from .hilfe import benutzer

    importiere(xlsx([z("A1", "x", 10**12, vh=10**12)]))
    client.force_login(benutzer("lager", KOMMISSIONIERUNG))
    assert client.get("/bestandsliste/").status_code == 200
    assert client.get("/").status_code == 200
